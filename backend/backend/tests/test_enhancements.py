import pytest
from pptx import Presentation
from pathlib import Path
from unittest.mock import patch, MagicMock

# Import all functions we are testing in this file
from worker.celery_app import (
    add_credits_to_slide, 
    remove_watermarks_from_masters, 
    generate_and_add_speaker_notes
)

# A directory for temporary test files (optional, but good practice)
TEST_OUTPUT_DIR = Path("test_outputs")
TEST_OUTPUT_DIR.mkdir(exist_ok=True)

@pytest.fixture
def sample_presentation():
    """Creates a simple, blank presentation with 3 slides for testing."""
    prs = Presentation()
    for _ in range(3):
        prs.slides.add_slide(prs.slide_layouts[6]) # Using a blank layout
    return prs

@pytest.fixture
def ppt_with_watermark():
    """Loads a pre-made presentation with a watermark on the master slide."""
    test_file_path = Path(__file__).parent / "test_watermark.pptx"
    assert test_file_path.exists(), "Test file 'test_watermark.pptx' not found in the tests/ directory."
    return Presentation(test_file_path)

@pytest.fixture
def slide_with_text():
    """Creates a presentation with a single slide containing some text."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content layout
    title = slide.shapes.title
    title.text = "Q3 Results"
    
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.text = "Sales are up 20% year-over-year."
    
    return slide

def test_add_credits_to_slide_with_custom_text(sample_presentation):
    """
    Tests that a hyperlinked text box can be added with custom text.
    """
    slide = sample_presentation.slides[0]
    slide_width = sample_presentation.slide_width
    slide_height = sample_presentation.slide_height
    
    custom_text = "My Custom Brand"
    custom_url = "https://mybrand.com"

    # Call the function with the new parameters
    add_credits_to_slide(slide, slide_width, slide_height, custom_text, custom_url)

    # Assertions
    assert len(slide.shapes) == 1
    new_shape = slide.shapes[0]
    assert new_shape.text_frame.text == custom_text
    
    hyperlink = new_shape.text_frame.paragraphs[0].runs[0].hyperlink
    assert hyperlink.address == custom_url

def test_remove_watermarks_from_masters(ppt_with_watermark):
    """
    Tests that a shape containing specific watermark text is removed from the slide master.
    """
    remove_watermarks_from_masters(ppt_with_watermark)
    slide_master = ppt_with_watermark.slide_masters[0]
    
    found_watermark = False
    for shape in slide_master.shapes:
        if shape.has_text_frame and "CONFIDENTIAL" in shape.text.upper():
            found_watermark = True
            break
            
    assert not found_watermark, "Watermark shape was not removed from the master slide"

def test_generate_speaker_notes(slide_with_text):
    """
    Tests that slide text is sent to the LLM and the response is added as notes.
    """
    fake_ai_response_text = "Key takeaway: Strong performance in Q3."

    # Patch the 'generate_content' method on the 'model' object directly.
    with patch('worker.celery_app.model.generate_content') as mock_generate_content:
        mock_response = MagicMock()
        mock_response.text = fake_ai_response_text
        mock_generate_content.return_value = mock_response

        # Call the function we intend to build
        generate_and_add_speaker_notes(slide_with_text)

        # 1. Assert that our mock method was called
        mock_generate_content.assert_called_once()
        call_args, _ = mock_generate_content.call_args
        prompt_sent_to_ai = call_args[0]
        assert "Q3 Results" in prompt_sent_to_ai
        assert "Sales are up 20%" in prompt_sent_to_ai

        # 2. Assert that the fake AI response was added to the notes slide
        notes_text = slide_with_text.notes_slide.notes_text_frame.text
        assert notes_text == fake_ai_response_text